# load_all_documents.py
# 加载所有三个课程的文档

import os
import json
import fitz  # PyMuPDF，用于解析 PDF
from docx import Document  # 用于解析 docx
from docx.opc.exceptions import PackageNotFoundError  # 捕获无效 docx 文件异常

# ===== 1. 配置路径 =====
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.dirname(os.path.dirname(SCRIPT_DIR))

# 原始教材默认放在项目内的 data/raw/textbooks；也可用环境变量覆盖。
TEXTBOOKS_DIR = os.environ.get(
    "RAG_TEXTBOOKS_DIR",
    os.path.join(BASE_DIR, "data", "raw", "textbooks"),
)
COURSE_DIRS = {
    "Computer Networks": os.path.join(TEXTBOOKS_DIR, "computer networks"),
    "Data Structures": os.path.join(TEXTBOOKS_DIR, "data structures"),
    "Operating System": os.path.join(TEXTBOOKS_DIR, "os"),
}

# 阶段二结构化文档输出目录
RAW_OUTPUT_DIR = os.path.join(BASE_DIR, "data", "raw_documents")
os.makedirs(RAW_OUTPUT_DIR, exist_ok=True)


# ===== 2. 工具函数：分别加载 txt / pdf / docx =====
def load_txt(file_path: str) -> str:
    """加载 TXT 文件内容"""
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception as e:
        print(f"读取 TXT 失败 {file_path}：{e}")
        return ""


def load_pdf(file_path: str) -> str:
    """加载 PDF 文件内容"""
    try:
        doc = fitz.open(file_path)
        texts = []
        for page in doc:
            texts.append(page.get_text("text"))
        return "\n".join(texts)
    except Exception as e:
        print(f"读取 PDF 失败 {file_path}：{e}")
        return ""


def load_docx(file_path: str) -> str:
    """加载 DOCX 文件内容（包含段落和表格）"""
    try:
        doc = Document(file_path)
        texts = []

        # 提取段落文本
        for para in doc.paragraphs:
            para_text = para.text.strip()
            if para_text:  # 跳过空行
                texts.append(para_text)

        # 提取表格内容（按行列格式化）
        for table in doc.tables:
            table_text = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                table_text.append("\t".join(row_cells))  # 单元格用制表符分隔
            if table_text:
                texts.append("\n".join(table_text))

        return "\n\n".join(texts)  # 段落/表格之间用空行分隔
    except PackageNotFoundError:
        print(f"无效的 DOCX 文件 {file_path}（可能是临时文件/损坏文件）")
        return ""
    except Exception as e:
        print(f"读取 DOCX 失败 {file_path}：{e}")
        return ""


# ===== 3. 主函数：遍历所有课程目录，转换为统一结构 =====
def load_all_courses_documents():
    all_documents = []
    course_counter = 0

    for course_name, input_dir in COURSE_DIRS.items():
        print(f"\n{'='*50}")
        print(f"正在处理课程: {course_name}")
        print(f"目录: {input_dir}")
        print(f"{'='*50}")

        # 检查目录是否存在
        if not os.path.exists(input_dir):
            print(f"警告：目录不存在 {input_dir}")
            continue

        # 统计当前课程的文档数量
        course_docs = []

        for filename in os.listdir(input_dir):
            # 过滤：跳过以 ~$ 开头的 Word 临时文件
            if filename.startswith('~$'):
                print(f"跳过 Word 临时文件: {filename}")
                continue

            file_path = os.path.join(input_dir, filename)
            text = ""

            # 根据后缀加载对应文件
            if filename.lower().endswith(".txt"):
                text = load_txt(file_path)
            elif filename.lower().endswith(".pdf"):
                text = load_pdf(file_path)
            elif filename.lower().endswith(".docx"):
                text = load_docx(file_path)
            else:
                # 其他格式暂时跳过
                print(f"跳过不支持的文件: {filename}")
                continue

            # 仅添加有有效内容的文档
            if text.strip():
                doc_id = f"{course_name}_{os.path.splitext(filename)[0]}"
                document = {
                    "id": doc_id,
                    "text": text,
                    "source": filename,
                    "subject": course_name
                }
                course_docs.append(document)
            else:
                print(f"跳过空内容文件: {filename}")

        # 将当前课程的文档添加到总列表
        all_documents.extend(course_docs)
        course_counter += 1
        print(f"课程 {course_name} 完成，共 {len(course_docs)} 篇文档")

    # 写入结构化 JSON 文件
    output_path = os.path.join(RAW_OUTPUT_DIR, "all_structured_documents.json")
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(all_documents, f, ensure_ascii=False, indent=2)

    print(f"\n{'='*50}")
    print(f"所有课程文档加载完成！")
    print(f"共处理 {course_counter} 个课程")
    print(f"总计 {len(all_documents)} 篇有效文档")
    print(f"已写入：{output_path}")
    print(f"{'='*50}")


if __name__ == "__main__":
    load_all_courses_documents()
